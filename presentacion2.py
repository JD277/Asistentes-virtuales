
# -*- coding: UTF-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# Crear una nueva presentacion
prs = Presentation()

# Definir los layouts que usaremos
title_slide_layout = prs.slide_layouts[0]
bullet_slide_layout = prs.slide_layouts[1]

# Diapositiva 1: Titulo
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Mecatronica: Integrando Disciplinas"
subtitle.text = "Una Introduccion Completa"

# Diapositiva 2: Definicion de Mecatronica
slide = prs.slides.add_slide(bullet_slide_layout)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
tf = body.text_frame

title.text = "Que es la Mecatronica?"
tf.text = "La mecatronica es un campo interdisciplinario de la ingenieria que combina:"
p = tf.add_paragraph()
p.text = "  - Ingenieria Mecanica"
p = tf.add_paragraph()
p.text = "  - Ingenieria Electronica"
p = tf.add_paragraph()
p.text = "  - Ingenieria de Control"
p = tf.add_paragraph()
p.text = "  - Ingenieria Informatica"

# Diapositiva 3: Componentes Clave
slide = prs.slides.add_slide(bullet_slide_layout)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
tf = body.text_frame

title.text = "Componentes Clave de la Mecatronica"
tf.text = "La mecatronica se basa en los siguientes componentes:"
p = tf.add_paragraph()
p.text = "  - Sensores y transductores: Para la adquisicion de datos."
p = tf.add_paragraph()
p.text = "  - Sistemas de control: Para el procesamiento de datos y la toma de decisiones."
p = tf.add_paragraph()
p.text = "  - Actuadores: Para la implementacion de acciones."
p = tf.add_paragraph()
p.text = "  - Software: Para la programacion y el control de sistemas."

# Diapositiva 4: Aplicaciones de la Mecatronica
slide = prs.slides.add_slide(bullet_slide_layout)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
tf = body.text_frame

title.text = "Aplicaciones de la Mecatronica"
tf.text = "La mecatronica tiene una amplia gama de aplicaciones, incluyendo:"
p = tf.add_paragraph()
p.text = "  - Robotica industrial: Automatizacion de procesos de fabricacion."
p = tf.add_paragraph()
p.text = "  - Sistemas automotrices: Control de vehiculos y seguridad."
p = tf.add_paragraph()
p.text = "  - Dispositivos medicos: Instrumentacion quirurgica y protesis."
p = tf.add_paragraph()
p.text = "  - Sistemas de automatizacion del hogar: Control de iluminacion y climatizacion."

# Diapositiva 5: Ventajas de la Mecatronica
slide = prs.slides.add_slide(bullet_slide_layout)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
tf = body.text_frame

title.text = "Ventajas de la Mecatronica"
tf.text = "La mecatronica ofrece varias ventajas:"
p = tf.add_paragraph()
p.text = "  - Mayor eficiencia: Optimiza el rendimiento de los sistemas."
p = tf.add_paragraph()
p.text = "  - Mayor precision: Permite un control preciso de los procesos."
p = tf.add_paragraph()
p.text = "  - Mayor flexibilidad: Adapta los sistemas a diferentes necesidades."
p = tf.add_paragraph()
p.text = "  - Mayor confiabilidad: Reduce el riesgo de fallas."

# Diapositiva 6: Desafios de la Mecatronica
slide = prs.slides.add_slide(bullet_slide_layout)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
tf = body.text_frame

title.text = "Desafios de la Mecatronica"
tf.text = "La mecatronica tambien presenta algunos desafios:"
p = tf.add_paragraph()
p.text = "  - Complejidad: Requiere conocimientos multidisciplinarios."
p = tf.add_paragraph()
p.text = "  - Costo: Puede ser mas costosa que las soluciones tradicionales."
p = tf.add_paragraph()
p.text = "  - Integracion: Requiere una integracion cuidadosa de los componentes."
p = tf.add_paragraph()
p.text = "  - Mantenimiento: Puede requerir personal especializado."

# Diapositiva 7: El futuro de la Mecatronica
slide = prs.slides.add_slide(bullet_slide_layout)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
tf = body.text_frame

title.text = "El Futuro de la Mecatronica"
tf.text = "El futuro de la mecatronica es prometedor:"
p = tf.add_paragraph()
p.text = "  - Inteligencia Artificial: Integracion de IA para sistemas mas inteligentes."
p = tf.add_paragraph()
p.text = "  - Internet de las Cosas (IoT): Conectividad y control remoto."
p = tf.add_paragraph()
p.text = "  - Robotica Avanzada: Robots mas autonomos y colaborativos."
p = tf.add_paragraph()
p.text = "  - Nanotecnologia: Miniaturizacion de sistemas mecatronicos."

# Diapositiva 8: Conclusion
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Conclusion"
subtitle.text = "La mecatronica es una disciplina clave para el futuro de la tecnologia y la innovacion."

# Guardar la presentacion
prs.save("presentacion_mecatronica.pptx")

#Abrir la presentacion con OS
os.startfile("presentacion_mecatronica.pptx")
